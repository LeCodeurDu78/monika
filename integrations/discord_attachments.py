"""Pièces jointes reçues sur Discord : enregistrement, lecture et mise en forme pour l'agent.

Traitement selon le type de fichier :
  - images (png, jpg, webp, gif)        -> analysées par le modèle vision (tools/vision/vision_tools.py) ;
  - PDF, DOCX, texte / code / CSV / JSON -> texte extrait ;
  - audio (dont messages vocaux Discord)  -> transcrit par Whisper (voice/voice_stt.py) ;
  - autre                                 -> simplement enregistré (Monika peut l'ouvrir avec ses outils).

Chaque fichier est gardé dans APP_DIR/discord_inbox (nettoyé après DISCORD_ATTACHMENT_RETENTION_DAYS jours) :
le chemin est donné à Monika, qui peut le relire avec read_file / run_script si l'extrait est tronqué.
Tout contenu extrait est encadré par <donnees_externes> : un document reçu peut contenir des instructions
piégées, ce sont des DONNÉES.

Les fonctions de ce module sont BLOQUANTES (vision, Whisper, lecture de PDF) : l'écouteur les appelle
hors de la boucle asyncio, via run_in_executor.
"""

import itertools
import re
import time
from dataclasses import dataclass
from pathlib import Path
from typing import Optional

from core.settings import settings
from core.untrusted import wrap_untrusted

INBOX_DIR = settings.APP_DIR / "discord_inbox"

IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".webp", ".gif"}
AUDIO_EXTENSIONS = {".ogg", ".oga", ".opus", ".mp3", ".wav", ".m4a", ".flac", ".aac"}
TEXT_EXTENSIONS = {
    ".txt", ".md", ".markdown", ".rst", ".csv", ".tsv", ".json", ".jsonl", ".yml", ".yaml", ".toml", ".ini",
    ".cfg", ".log", ".xml", ".html", ".htm", ".sql", ".sh", ".py", ".js", ".ts", ".tsx", ".jsx", ".java",
    ".c", ".h", ".cpp", ".hpp", ".go", ".rs",
}
PDF_MAX_PAGES = 80

_counter = itertools.count(1)
_cleaned = False


@dataclass
class Attachment:
    """Pièce jointe déjà téléchargée."""

    filename: str
    data: bytes
    content_type: str = ""


def _safe_name(filename: str) -> str:
    """Nom de fichier sûr (pas de chemin, pas de caractères spéciaux), horodaté pour éviter les collisions."""
    base = Path(str(filename or "fichier")).name
    base = re.sub(r"[^\w.\-]", "_", base).strip("._") or "fichier"
    stem, dot, ext = base.rpartition(".")
    if not dot:
        stem, ext = base, ""
    stem = stem[:60] or "fichier"
    ext = f".{ext[:10]}" if ext else ""
    return f"{int(time.time())}_{next(_counter)}_{stem}{ext}"


def _cleanup_inbox() -> None:
    """Supprime les vieux fichiers reçus (une fois par session)."""
    global _cleaned
    if _cleaned or not INBOX_DIR.exists():
        return
    _cleaned = True
    cutoff = time.time() - settings.DISCORD_ATTACHMENT_RETENTION_DAYS * 86400
    for path in INBOX_DIR.iterdir():
        try:
            if path.is_file() and path.stat().st_mtime < cutoff:
                path.unlink()
        except OSError:
            pass


def _save(attachment: Attachment) -> Path:
    INBOX_DIR.mkdir(parents=True, exist_ok=True)
    _cleanup_inbox()
    path = INBOX_DIR / _safe_name(attachment.filename)
    path.write_bytes(attachment.data)
    return path


def _kind(attachment: Attachment) -> str:
    ext = Path(attachment.filename).suffix.lower()
    ctype = (attachment.content_type or "").split(";")[0].strip().lower()
    if ext in IMAGE_EXTENSIONS or ctype.startswith("image/"):
        return "image"
    if ext in AUDIO_EXTENSIONS or ctype.startswith("audio/"):
        return "audio"
    if ext == ".pdf" or ctype == "application/pdf":
        return "pdf"
    if ext == ".docx":
        return "docx"
    if ext in TEXT_EXTENSIONS or ctype.startswith("text/"):
        return "text"
    return "other"


def _human_size(n: int) -> str:
    if n < 1024:
        return f"{n} o"
    if n < 1024 * 1024:
        return f"{n / 1024:.0f} Ko"
    return f"{n / (1024 * 1024):.1f} Mo"


# --- extraction par type ------------------------------------------------------------------------
def _describe_image(path: Path, content_type: str, user_text: str) -> str:
    from tools.vision.vision_tools import analyze_image  # import tardif : charge le client vision

    if user_text:
        prompt = (
            f"L'utilisateur envoie cette image avec le message : « {user_text} ». Décris-la en détail, "
            "transcris le texte visible, et réponds à ce qu'il demande si c'est une question sur l'image."
        )
    else:
        prompt = "Décris cette image en détail et transcris le texte visible."
    mime = content_type.split(";")[0].strip() if (content_type or "").lower().startswith("image/") else None
    return analyze_image(str(path), prompt, mime_type=mime)


def _extract_pdf(path: Path) -> str:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    if reader.is_encrypted:
        try:
            if not reader.decrypt(""):
                return "(PDF protégé par un mot de passe : lecture impossible)"
        except Exception:
            return "(PDF protégé par un mot de passe : lecture impossible)"
    limit = settings.DISCORD_ATTACHMENT_MAX_CHARS
    parts: list[str] = []
    total = 0
    pages = reader.pages
    for number, page in enumerate(pages[:PDF_MAX_PAGES], 1):
        text = (page.extract_text() or "").strip()
        if text:
            parts.append(f"[page {number}]\n{text}")
            total += len(text)
        if total > limit * 2:  # inutile de lire tout un gros PDF pour n'en garder que le début
            break
    if not parts:
        return "(aucun texte extractible : PDF probablement scanné, essaie de l'envoyer en image)"
    extra = f"\n[... {len(pages) - PDF_MAX_PAGES} pages non lues]" if len(pages) > PDF_MAX_PAGES else ""
    return "\n\n".join(parts) + extra


def _extract_docx(path: Path) -> str:
    from docx import Document

    document = Document(str(path))
    lines = [p.text for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            lines.append(" | ".join(cell.text.strip() for cell in row.cells))
    return "\n".join(lines) or "(document vide)"


def _extract_text(data: bytes) -> str:
    return data.decode("utf-8", errors="replace")


def _transcribe(path: Path) -> str:
    from voice.voice_stt import transcribe_file  # import tardif : charge Whisper à la première demande

    text = transcribe_file(str(path))
    return text or "(aucune parole détectée)"


def _extract(kind: str, path: Path, attachment: Attachment, user_text: str) -> Optional[str]:
    if kind == "image":
        return _describe_image(path, attachment.content_type, user_text)
    if kind == "pdf":
        return _extract_pdf(path)
    if kind == "docx":
        return _extract_docx(path)
    if kind == "text":
        return _extract_text(attachment.data)
    if kind == "audio":
        return _transcribe(path)
    return None


_KIND_LABEL = {
    "image": "image (analyse vision)",
    "pdf": "document PDF (texte extrait)",
    "docx": "document Word (texte extrait)",
    "text": "fichier texte",
    "audio": "audio (transcription)",
    "other": "fichier",
}


def _clip(text: str, path: Path) -> str:
    limit = settings.DISCORD_ATTACHMENT_MAX_CHARS
    if len(text) <= limit:
        return text
    return (
        f"{text[:limit]}\n[... tronqué : {len(text)} caractères au total ; "
        f"fichier complet : {path} — relis-le avec read_file ou run_script]"
    )


def describe_attachment(attachment: Attachment, index: int, total: int, user_text: str) -> str:
    """Bloc de texte décrivant UNE pièce jointe pour l'agent (ne lève jamais d'exception)."""
    kind = _kind(attachment)
    shown_name = Path(str(attachment.filename or "fichier")).name  # jamais de chemin dans ce qu'on affiche
    header = f"[Pièce jointe {index}/{total} : {shown_name} — {_KIND_LABEL[kind]}, {_human_size(len(attachment.data))}"
    try:
        path = _save(attachment)
    except OSError as exc:
        return f"{header}] Impossible d'enregistrer le fichier ({type(exc).__name__}: {exc})."

    header += f", enregistrée dans {path}]"
    try:
        extracted = _extract(kind, path, attachment, user_text)
    except Exception as exc:  # une pièce jointe illisible ne doit pas empêcher de répondre au reste
        print(f"⚠️ [Discord] Lecture de {shown_name} impossible : {type(exc).__name__}: {exc}")
        return f"{header}\nContenu illisible ({type(exc).__name__}: {exc})."

    if extracted is None:
        return f"{header}\nType de fichier non lu automatiquement : ouvre-le avec tes outils si besoin."
    return f"{header}\n{wrap_untrusted(_clip(extracted, path), f'pièce jointe {shown_name}')}"


def build_message(user_text: str, attachments: list[Attachment], skipped: Optional[list[str]] = None) -> str:
    """Message utilisateur complet : son texte suivi d'un bloc par pièce jointe (et des fichiers ignorés)."""
    blocks = [describe_attachment(att, i, len(attachments), user_text) for i, att in enumerate(attachments, 1)]
    blocks.extend(f"[Pièce jointe ignorée : {reason}]" for reason in (skipped or []))
    return "\n\n".join(part for part in [user_text.strip(), *blocks] if part)
